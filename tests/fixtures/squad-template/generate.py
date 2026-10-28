"""Deterministic fictional workbooks matching the supplied official-template contract.
Not replacements for any real-photo OCR fixture. Requires openpyxl to regenerate.
"""
import re
from pathlib import Path
from datetime import datetime
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font
ROOT=Path(__file__).parent
HEAD=['Player Name *']
LEGACY=['Player Name *','Default No.','Preferred Position','Second Position','Status']
NAMES=['Aria Alder','Bella Birch','Cora Cedar','Dara Dale','Ella Elm','Fia Field','Grace Grove','Hana Hazel','Ivy Iris','Jade Juniper','Kara Lake','Lena Linden','Mara Maple','Nessa Oak','Orla Pine','Pia Reed','Rosa Rowan','Sara Sage','Tara Stone','Una Vale','Vera Willow','Willa Wood','Yara Ash','Zara Bay','Ada Brook','Bea Fern','Cara Hill','Dee Moss','Eva Park','Faye Rye']
def write(name,rows,missing=False,headers=HEAD):
 w=Workbook();s=w.active;s.title='Squad Setup'
 s.append(['Official Sideline Squad Import Template'])
 for row in [('Club Name','Fictional County'),('Squad Name','U16 Girls'),('Season','2027'),('Code','LGFA'),('Coach / Manager','Test Coach')]:s.append(row)
 if not missing:
  s=w.create_sheet('Players');s.append(headers)
  for row in rows:s.append(row)
  for i in range(len(rows),40):
   s.append([None]*len(headers))
   for c in s[s.max_row]:c.font=Font(name='Calibri')
 w.properties.created=w.properties.modified=datetime(2026,1,1)
 data=BytesIO();w.save(data)
 with ZipFile(BytesIO(data.getvalue())) as source,ZipFile(ROOT/name,'w',ZIP_DEFLATED) as target:
  for entry in sorted(source.namelist()):
   info=ZipInfo(entry,(2026,1,1,0,0,0));info.compress_type=ZIP_DEFLATED;payload=source.read(entry)
   if entry=="docProps/core.xml":payload=re.sub(rb"(<dcterms:modified[^>]*>).*?(</dcterms:modified>)",rb"\g<1>2026-01-01T00:00:00Z\g<2>",payload)
   target.writestr(info,payload)
write('County-30.xlsx',[[n] for n in NAMES])
write('County-29.xlsx',[[n] for n in NAMES[:29]])
write('Legacy-rich.xlsx',[[NAMES[0],23,'Not a position'*20,'Ignored','Injured'],[NAMES[1],None,'Midfield','Full Back','Inactive'],[None,None,'Ignored','Ignored','Unavailable']],headers=LEGACY)
write('Unicode-names.xlsx', [[' \nÉabha Ní Bhriain\r\n'],['Siún\n  O’Neill'],['mÁire McDONAGH'],['Same Name'],['Same Name']])
write('Name-only.xlsx',[[NAMES[0]],[NAMES[1]]],headers=['Player Name *'])
write('Club-21.xlsx',[[n] for n in NAMES[:21]])
write('Review-errors.xlsx', [['Same Name',6],['Same Name',6],['Third Player','oops'],['',8]], headers=['Player Name *','Default No.'])
write('Missing-Players.xlsx',[],missing=True)
write('Wrong-columns.xlsx',[[NAMES[0]]],headers=['Name','Number','Position','Other','State'])
(ROOT/'Malformed.xlsx').write_bytes(b'This is not an Excel workbook.')
