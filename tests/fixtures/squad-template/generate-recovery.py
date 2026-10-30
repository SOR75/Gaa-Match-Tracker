"""Representative official workbook with metadata/instructions before Players.
Fictional names; this is a spreadsheet-routing fixture, never a photo substitute.
"""
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
w=Workbook();s=w.active;s.title='Squad Setup'
for row in [('Club Name','Offaly LGFA'),('Squad Name','U15 Development'),('Season','2027'),('Code','LGFA'),('Coach / Manager','SHANE'),('Complete the squad details...',),('Player Name only...',),('Default No. is ignored',),('Sideline will create permanent IDs',)]:s.append(row)
s=w.create_sheet('Instructions');s.append(['Complete the squad details before uploading.'])
s=w.create_sheet('Players');s.append(['Official Sideline Squad Template']);s.append(['Player Name only: add names below the column heading.']);s.append(['Player Name *','Default No.'])
for row in [('  Éabha Ní Bhriain\n',23),('Test Player Two',None),('Test Player Two',8),('',None)]:s.append(row)
s.append(["Sideline will create permanent IDs when confirmed."])
s.append(["Player Name only: leave unused rows empty."])
w.properties.created=w.properties.modified=datetime(2026,1,1)
w.save(Path(__file__).with_name('Official-routing.xlsx'))
